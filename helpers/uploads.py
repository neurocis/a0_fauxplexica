"""Uploaded-file RAG adapter for A0_Fauxplexica.

This module intentionally reuses the `_memory` plugin's FAISS/embedding
infrastructure.  Each A0 chat context gets a dedicated `_memory` subdirectory:
`fauxplexica_uploads/<context_id>/`, so uploaded-file chunks do not pollute the
agent's normal long-term memories.
"""

from __future__ import annotations

import asyncio
import logging
import mimetypes
import uuid
from pathlib import Path
from typing import Any

try:
    from langchain_core.documents import Document
except ImportError:  # pragma: no cover - lightweight fallback for minimal test envs
    class Document:  # type: ignore[no-redef]
        def __init__(self, page_content: str, metadata: dict | None = None):
            self.page_content = page_content
            self.metadata = metadata or {}

from .utils import split_text

logger = logging.getLogger(__name__)


async def get_uploads_memory(memory_subdir: str, log_item: Any | None = None) -> Any:
    """Return a `_memory` wrapper for an uploads namespace.

    Kept as a small seam so tests can mock `_memory` without importing the
    heavyweight FAISS/LangChain stack. Runtime behavior still delegates to
    `_memory.helpers.memory.Memory.get_by_subdir`.
    """
    from plugins._memory.helpers.memory import Memory

    return await Memory.get_by_subdir(
        memory_subdir,
        log_item=log_item,
        preload_knowledge=False,
    )


_SUPPORTED_TEXT_EXTS = {".txt", ".md", ".markdown", ".csv"}
_SUPPORTED_DOCX_EXTS = {".docx"}
_SUPPORTED_PDF_EXTS = {".pdf"}
_DEFERRED_IMAGE_MIMES = {"image/png", "image/jpeg", "image/gif", "image/webp", "image/bmp", "image/tiff"}


class UploadsManager:
    """Thin per-context upload-RAG adapter over `_memory`'s vector store.

    Args:
        agent_context: An A0 ``AgentContext``. The manager binds to
            ``fauxplexica_uploads/<context_id>`` via ``Memory.get_by_subdir``.
    """

    def __init__(self, agent_context: Any):
        self.context = agent_context
        self.agent = agent_context.get_agent() if hasattr(agent_context, "get_agent") else getattr(agent_context, "agent0", None)
        self.context_id = str(getattr(agent_context, "id", "default"))
        self.memory_subdir = f"fauxplexica_uploads/{self.context_id}"
        self._memory: Any | None = None
        self.last_ingest_result: dict[str, Any] | None = None

    async def _get_memory(self) -> Any:
        """Return the scoped `_memory` wrapper for this chat context."""
        if self._memory is None:
            log_item = None
            if self.agent is not None:
                try:
                    log_item = self.agent.context.log.log(
                        type="util",
                        heading=f"Initializing Fauxplexica uploads VectorDB in '/{self.memory_subdir}'",
                    )
                except Exception:  # pragma: no cover - logging must never block ingestion
                    log_item = None
            self._memory = await get_uploads_memory(self.memory_subdir, log_item=log_item)
        return self._memory

    async def ingest_file(self, path: str, *, file_id: str | None = None) -> str | dict[str, str]:
        """Extract, chunk, embed, and index one file.

        Supported in Phase 1:
        - PDF via ``pdfplumber``
        - TXT/MD/CSV via UTF-8 text read
        - DOCX via ``python-docx``

        Images and unknown/binary types are deferred. Per-file failures are
        returned as ``{"error": "..."}`` and logged rather than raised.
        """
        file_path = Path(path).expanduser()
        assigned_file_id = file_id or uuid.uuid4().hex
        self.last_ingest_result = None

        try:
            if not file_path.exists() or not file_path.is_file():
                return self._error(f"File not found: {path}")

            content, detected_type = await self._extract_text(file_path)
            if not content.strip():
                return self._error(f"No extractable text found in {file_path.name}")

            chunks = [chunk for chunk in split_text(content, 4000, 500) if chunk.strip()]
            if not chunks:
                return self._error(f"No chunks produced for {file_path.name}")

            total_chunks = len(chunks)
            docs = [
                Document(
                    page_content=chunk,
                    metadata={
                        "area": "fauxplexica_uploads",
                        "file_id": assigned_file_id,
                        "source_path": str(file_path),
                        "filename": file_path.name,
                        "chunk_index": i,
                        "total_chunks": total_chunks,
                        "mime_type": detected_type,
                        "namespace": self.memory_subdir,
                        "url": f"file_id://{assigned_file_id}",
                    },
                )
                for i, chunk in enumerate(chunks)
            ]

            memory = await self._get_memory()
            ids = await memory.insert_documents(docs)
            self.last_ingest_result = {
                "file_id": assigned_file_id,
                "chunks_indexed": len(ids),
                "filename": file_path.name,
                "source_path": str(file_path),
                "mime_type": detected_type,
            }
            return assigned_file_id
        except Exception as exc:  # defensive per Recon §8: do not let one file kill pipeline
            logger.exception("Failed to ingest upload %s", path)
            return self._error(str(exc))

    RRF_K: ClassVar[int] = 60

    async def search(
        self,
        query: str | list[str],
        file_ids: list[str] | None = None,
        *,
        k: int = 8,
    ) -> list[dict[str, Any]]:
        """Search uploaded-file chunks for ``query``.

        ``query`` may be a single string or a list of strings. When multiple
        queries are supplied, results are fused with Reciprocal Rank Fusion
        (``RRF_K = 60``) to match Vane's :class:`UploadStore` semantics.

        If ``file_ids`` is provided and non-empty, it is treated as an explicit
        override by callers and used as a metadata filter. Returned items use
        ``{content, metadata, score, rrf_score?}``; `_memory` exposes
        normalized relevance scores when available, otherwise ``score`` is
        ``None``.
        """
        queries = [str(q).strip() for q in (query if isinstance(query, (list, tuple, set)) else [query]) if str(q).strip()]
        if not queries:
            return []
        try:
            memory = await self._get_memory()
            filter_expr = "area == 'fauxplexica_uploads'"
            if file_ids:
                safe_ids = ", ".join(repr(str(fid)) for fid in file_ids)
                filter_expr += f" and file_id in [{safe_ids}]"
            top_k = max(1, int(k))

            per_query_results: list[list[dict[str, Any]]] = []
            for q in queries:
                hits = await memory.search_similarity_threshold(
                    q,
                    limit=top_k,
                    threshold=0.0,
                    filter=filter_expr,
                )
                per_query_results.append([self._format_search_result(item) for item in hits])

            if len(per_query_results) == 1:
                return per_query_results[0]
            return self._reciprocal_rank_fuse(per_query_results, top_k=top_k)
        except Exception as exc:
            logger.exception("Failed to search uploads")
            return [{"error": str(exc), "content": "", "metadata": {}, "score": None}]

    @classmethod
    def _reciprocal_rank_fuse(
        cls,
        ranked_lists: list[list[dict[str, Any]]],
        *,
        top_k: int,
    ) -> list[dict[str, Any]]:
        """Fuse multiple ranked result lists via Reciprocal Rank Fusion."""
        fused: dict[str, dict[str, Any]] = {}
        for ranked in ranked_lists:
            for rank, item in enumerate(ranked):
                meta = item.get("metadata") or {}
                key = str(
                    meta.get("id")
                    or meta.get("url")
                    or f"{meta.get('file_id', '')}::{meta.get('chunk_index', '')}::{item.get('content', '')[:48]}"
                )
                contribution = 1.0 / (rank + 1 + cls.RRF_K)
                if key in fused:
                    fused[key]["rrf_score"] = float(fused[key].get("rrf_score", 0.0)) + contribution
                else:
                    clone = dict(item)
                    clone["metadata"] = dict(meta)
                    clone["rrf_score"] = contribution
                    fused[key] = clone
        ordered = sorted(fused.values(), key=lambda entry: entry.get("rrf_score", 0.0), reverse=True)
        return ordered[:top_k]

    async def list_files(self) -> list[dict[str, Any]]:
        """List distinct uploaded files indexed in this context namespace."""
        try:
            memory = await self._get_memory()
            docs = self._all_upload_docs(memory)
            grouped: dict[str, dict[str, Any]] = {}
            for doc in docs:
                meta = dict(doc.metadata or {})
                fid = str(meta.get("file_id") or "")
                if not fid:
                    continue
                entry = grouped.setdefault(
                    fid,
                    {
                        "file_id": fid,
                        "filename": meta.get("filename", ""),
                        "source_path": meta.get("source_path", ""),
                        "mime_type": meta.get("mime_type", ""),
                        "chunk_count": 0,
                        "total_chunks": meta.get("total_chunks"),
                    },
                )
                entry["chunk_count"] += 1
                if meta.get("total_chunks") is not None:
                    entry["total_chunks"] = meta.get("total_chunks")
            return sorted(grouped.values(), key=lambda item: (item.get("filename") or "", item["file_id"]))
        except Exception as exc:
            logger.exception("Failed to list uploads")
            return [{"error": str(exc)}]

    async def remove_file(self, file_id: str) -> int:
        """Delete all chunks for ``file_id`` from this context namespace."""
        try:
            memory = await self._get_memory()
            ids = [
                doc.metadata["id"]
                for doc in self._all_upload_docs(memory)
                if doc.metadata.get("file_id") == file_id and doc.metadata.get("id")
            ]
            if not ids:
                return 0
            removed = await memory.delete_documents_by_ids(ids)
            return len(removed)
        except Exception:
            logger.exception("Failed to remove upload file_id=%s", file_id)
            return 0

    async def _extract_text(self, path: Path) -> tuple[str, str]:
        mime_type, _encoding = mimetypes.guess_type(str(path))
        ext = path.suffix.lower()
        detected = mime_type or self._mime_from_extension(ext)

        if detected in _DEFERRED_IMAGE_MIMES or (mime_type or "").startswith("image/"):
            raise ValueError("Image uploads are deferred to Phase 2")
        if ext in _SUPPORTED_PDF_EXTS or detected == "application/pdf":
            return await asyncio.to_thread(self._extract_pdf, path), "application/pdf"
        if ext in _SUPPORTED_DOCX_EXTS or detected == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return await asyncio.to_thread(self._extract_docx, path), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        if ext in _SUPPORTED_TEXT_EXTS or (detected or "").startswith("text/"):
            return await asyncio.to_thread(path.read_text, encoding="utf-8", errors="replace"), detected or "text/plain"
        raise ValueError(f"Unsupported Phase 1 upload type: {detected or ext or 'unknown'}")

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        import pdfplumber

        pages: list[str] = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
        return "\n\n".join(pages)

    @staticmethod
    def _extract_docx(path: Path) -> str:
        import docx

        document = docx.Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    @staticmethod
    def _mime_from_extension(ext: str) -> str:
        return {
            ".txt": "text/plain",
            ".md": "text/markdown",
            ".markdown": "text/markdown",
            ".csv": "text/csv",
            ".pdf": "application/pdf",
            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".png": "image/png",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".gif": "image/gif",
            ".webp": "image/webp",
        }.get(ext, "application/octet-stream")

    @staticmethod
    def _format_search_result(item: Any) -> dict[str, Any]:
        score = None
        doc = item
        if isinstance(item, tuple) and item:
            doc = item[0]
            if len(item) > 1:
                score = item[1]
        return {
            "content": getattr(doc, "page_content", ""),
            "metadata": dict(getattr(doc, "metadata", {}) or {}),
            "score": score,
        }

    @staticmethod
    def _all_upload_docs(memory: Any) -> list[Document]:
        raw_docs = memory.db.get_all_docs()
        docs = list(raw_docs.values()) if isinstance(raw_docs, dict) else list(raw_docs)
        return [doc for doc in docs if getattr(doc, "metadata", {}).get("area") == "fauxplexica_uploads"]

    def _error(self, message: str) -> dict[str, str]:
        logger.error("Fauxplexica upload error: %s", message)
        self.last_ingest_result = {"error": message}
        return {"error": message}


__all__ = ["Document", "UploadsManager", "get_uploads_memory"]
